import streamlit as st
import requests
from docx import Document
from reportlab.pdfgen import canvas
st.set_page_config(
    page_title="LegalEase AI",
    page_icon="⚖️",
    layout="centered"
)

st.title("⚖️ LegalEase AI")
st.write("AI-Powered Legal Document Generator")

document_type = st.text_input(
    "Document Type",
    placeholder="Example: Employment Contract"
)

parties = st.text_area(
    "Parties Involved",
    placeholder="Example: John Doe (Employee), ABC Corp (Employer)"
)

terms = st.text_area(
    "Terms & Conditions",
    placeholder="Enter terms separated by semicolons (;)"
)

dates = st.text_input(
    "Effective Date",
    placeholder="Example: 10/04/2025"
)

if st.button("Generate Document"):

    if not document_type or not parties or not terms or not dates:
        st.warning("Please fill in all fields.")

    else:
        data = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "dates": dates
        }

        try:
            response = requests.post(
                "http://127.0.0.1:8000/generate",
                json=data
            )

            if response.status_code == 200:

                result = response.json().get("document")

                if result and isinstance(result, str):

                    st.success("Document Generated Successfully!")

                    edited_document = st.text_area(
                        "Edit Document",
                        value=result,
                        height=400
                    )
                                        # Create DOCX
                    doc = Document()
                    doc.add_heading("LegalEase AI Document", level=1)

                    for line in edited_document.split("\n"):
                        doc.add_paragraph(line)

                    docx_path = "LegalEase_Document.docx"
                    doc.save(docx_path)

                    # Create PDF
                    pdf_path = "LegalEase_Document.pdf"
                    pdf = canvas.Canvas(pdf_path)

                    y = 800

                    for line in edited_document.split("\n"):
                        if y < 50:
                            pdf.showPage()
                            y = 800

                        pdf.drawString(50, y, line[:100])
                        y -= 15

                    pdf.save()

                    st.download_button(
                        "Download TXT",
                        data=edited_document,
                        file_name="LegalEase_Document.txt",
                        mime="text/plain"
                    )
                    
                    with open(docx_path, "rb") as file:
                        st.download_button(
                            "Download DOCX",
                            data=file,
                            file_name="LegalEase_Document.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )

                    with open(pdf_path, "rb") as file:
                        st.download_button(
                            "Download PDF",
                            data=file,
                            file_name="LegalEase_Document.pdf",
                            mime="application/pdf"
                        )
                        

                else:
                    st.error(
                        "Document was not returned by the AI. "
                        "Please try again."
                    )

            elif response.status_code == 429:

                st.warning(
                    "⚠️ Gemini API quota exceeded.\n\n"
                    "Please try again after the quota resets."
                )

            else:

                try:
                    error_detail = response.json().get(
                        "detail",
                        "An error occurred."
                    )
                except:
                    error_detail = "An error occurred."

                st.error(error_detail)

        except Exception as e:
            st.error(f"Connection Error: {e}")